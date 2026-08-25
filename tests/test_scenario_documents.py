from __future__ import annotations

from io import BytesIO
from pathlib import Path

import httpx
import pytest
from docx import Document

from pytho_arena.scenario.documents import export_scenario_document, import_scenario_document
from pytho_arena.scenario.models import load_scenario_file
from pytho_arena.server.app import create_app


def background_brief() -> bytes:
    document = Document()
    document.add_heading("Background", level=1)
    document.add_paragraph("Small unmanned systems create an optical surveillance challenge.")
    document.add_heading("Scope", level=1)
    document.add_paragraph("This effort will focus on Harbor Test Installation.")
    document.add_heading("Relevant Authorities", level=1)
    document.add_paragraph("Civil aviation and installation security coordinate the exercise.")
    document.add_heading("Base Overview and Vulnerabilities", level=1)
    document.add_paragraph("The coastal and urban site contains gates and generic utilities.")
    document.add_heading("Threat Surveillance Capabilities", level=1)
    document.add_paragraph("Commercial small UAS may approach from public or maritime areas.")
    document.add_heading("Blue Surveillance and Tracking", level=1)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Modality"
    table.rows[0].cells[1].text = "Active/Passive"
    for modality, operation in [("Radar", "Active"), ("EO", "Passive"), ("RF", "Passive")]:
        cells = table.add_row().cells
        cells[0].text = modality
        cells[1].text = operation
    document.add_heading("Blue Effectors", level=1)
    document.add_paragraph("Control jamming and directed energy are considered abstractly.")
    document.add_heading("Detection and Vulnerability Views", level=1)
    document.add_paragraph("Coverage views support exercise review.")
    document.add_heading("References", level=1)
    document.add_paragraph("Synthetic test reference.", style="List Number")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_document_import_builds_reviewed_coastal_scene():
    result = import_scenario_document(background_brief(), "urban-brief.docx")

    assert result.scenario.name == "Harbor Test Installation — Fictionalized C-UAS"
    assert result.scenario.visualization is not None
    assert result.scenario.visualization.environment == "coastal_urban"
    assert len(result.scenario.visualization.static_assets) >= 20
    assert result.report.sensor_modalities == ["Radar", "EO", "RF"]
    assert len(result.report.asset_resolution) == 5
    assert any("maps" in item for item in result.report.safety_transformations)
    assert result.scenario.actor_groups[0].attack_probability == 0.0


def test_document_import_rejects_non_docx():
    with pytest.raises(ValueError, match="valid DOCX"):
        import_scenario_document(b"not a zip package", "bad.docx")


def test_document_export_matches_background_brief_contract(tmp_path: Path):
    scenario = load_scenario_file(
        Path("src/pytho_arena/scenario/library/strait-sentinel-airbase.json")
    )
    output = export_scenario_document(scenario, tmp_path / "raven-background.docx")
    rendered = Document(output)
    headings = [
        paragraph.text for paragraph in rendered.paragraphs if paragraph.style.name == "Heading 1"
    ]

    assert headings == [
        "Background",
        "Scope",
        "Relevant Authorities and Coordination",
        "Base Overview and Protected Assets",
        "Threat Surveillance Capabilities",
        "Blue Surveillance and Tracking",
        "Blue Effectors",
        "Detection and Synthetic Performance Views",
        "References",
        "Traceability",
    ]
    assert len(rendered.tables) == 3
    assert "synthetic" in " ".join(paragraph.text.lower() for paragraph in rendered.paragraphs)
    assert output.stat().st_size > 10_000


@pytest.mark.asyncio
async def test_api_imports_docx_and_exports_saved_scenario(tmp_path: Path):
    app = create_app(tmp_path / "data")
    async with app.router.lifespan_context(app):
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=app), base_url="http://test"
        ) as client:
            imported = await client.post(
                "/api/scenarios/import-document",
                files={
                    "file": (
                        "urban-brief.docx",
                        background_brief(),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
            )
            assert imported.status_code == 201, imported.text
            payload = imported.json()
            scenario = payload["record"]["scenario"]
            assert scenario["visualization"]["environment"] == "coastal_urban"
            assert payload["report"]["scenarioVersion"] == 1

            exported = await client.get(
                f"/api/scenarios/{scenario['scenario_id']}/background-document?version=1"
            )
            assert exported.status_code == 200
            assert exported.headers["content-type"].startswith(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
            assert Document(BytesIO(exported.content)).paragraphs[0].text.startswith("FICTIONAL")
