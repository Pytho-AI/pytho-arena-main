"""Bidirectional scenario documents for the JLC-style background brief contract."""

from __future__ import annotations

import hashlib
import re
from collections import defaultdict
from io import BytesIO
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

from pydantic import BaseModel, ConfigDict

from .models import ScenarioSpec

MAX_DOCX_BYTES = 25 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024

SECTION_ALIASES = {
    "background": "background",
    "scope": "scope",
    "relevant authorities": "coordination",
    "authorities": "coordination",
    "coordination": "coordination",
    "base overview and vulnerabilities": "protected_assets",
    "base overview": "protected_assets",
    "protected assets": "protected_assets",
    "threat surveillance capabilities": "threat",
    "threat capabilities": "threat",
    "blue surveillance and tracking": "sensors",
    "sensors": "sensors",
    "blue effectors": "responses",
    "responses": "responses",
    "detection and vulnerability views": "coverage",
    "coverage views": "coverage",
    "references": "references",
}


class AssetResolution(BaseModel):
    model_config = ConfigDict(extra="forbid")

    source_category: str
    scene_asset: str
    implementation: str
    status: str = "resolved"


class ScenarioDocumentImportReport(BaseModel):
    model_config = ConfigDict(extra="forbid")

    source_filename: str
    source_sha256: str
    source_size_bytes: int
    source_site_name: str | None = None
    matched_sections: list[str]
    sensor_modalities: list[str]
    response_modalities: list[str]
    embedded_images: int
    asset_resolution: list[AssetResolution]
    safety_transformations: list[str]
    assumptions: list[str]
    warnings: list[str]


class ScenarioDocumentImport(BaseModel):
    model_config = ConfigDict(extra="forbid")

    scenario: ScenarioSpec
    report: ScenarioDocumentImportReport


def _document_class():
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - exercised by install smoke tests
        raise RuntimeError(
            "Scenario document support requires the server extra: pip install 'pytho-arena[server]'"
        ) from exc
    return Document


def _validate_docx(data: bytes) -> int:
    if not data or len(data) > MAX_DOCX_BYTES:
        raise ValueError("DOCX must be non-empty and no larger than 25 MB")
    try:
        with ZipFile(BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > 1_000:
                raise ValueError("DOCX contains too many package parts")
            if sum(member.file_size for member in members) > MAX_UNCOMPRESSED_BYTES:
                raise ValueError("DOCX expands beyond the 100 MB safety limit")
            if "word/document.xml" not in archive.namelist():
                raise ValueError("file is not a Word DOCX document")
            return sum(name.startswith("word/media/") for name in archive.namelist())
    except BadZipFile as exc:
        raise ValueError("file is not a valid DOCX package") from exc


def _normalize_heading(value: str) -> str:
    return re.sub(r"\s+", " ", value.strip().lower()).rstrip(":")


def _extract_sections(document: Any) -> dict[str, list[str]]:
    sections: dict[str, list[str]] = defaultdict(list)
    current = "background"
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        normalized = _normalize_heading(text)
        slot = SECTION_ALIASES.get(normalized)
        is_heading = paragraph.style.name.startswith("Heading") or slot is not None
        if slot is not None and is_heading:
            current = slot
            sections.setdefault(current, [])
            continue
        sections[current].append(text)
    return dict(sections)


def _extract_modalities(document: Any) -> list[str]:
    for table in document.tables:
        if not table.rows or not table.rows[0].cells:
            continue
        if "modality" not in table.rows[0].cells[0].text.strip().lower():
            continue
        modalities: list[str] = []
        for row in table.rows[1:]:
            value = row.cells[0].text.strip()
            if not value or value.lower() == "modality" or value in modalities:
                continue
            modalities.append(value)
        return modalities[:6]
    return []


def _extract_site_name(scope: str) -> str | None:
    match = re.search(r"(?:focus on|focused on)\s+(.+?)(?:\.\s|;|$)", scope, re.IGNORECASE)
    if not match:
        return None
    value = re.sub(r"\s+", " ", match.group(1)).strip(" .")
    return value[:100] or None


def _response_modalities(text: str) -> list[str]:
    candidates = [
        ("Control jamming", r"control jamming"),
        ("Sensor jamming", r"sensor jamming"),
        ("Directed energy", r"directed energy|high-power microwave|laser"),
        ("Kinetic", r"kinetic|projectile|anti-drone"),
    ]
    return [label for label, pattern in candidates if re.search(pattern, text, re.IGNORECASE)]


def _id(value: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")
    if not slug or not slug[0].isalpha():
        slug = f"item-{slug}"
    if len(slug) < 3:
        slug = f"{slug}-sensor"
    return slug[:63]


def _static_assets() -> list[dict[str, Any]]:
    assets: list[dict[str, Any]] = [
        {"id": "operations-center", "kind": "office_block", "position": [80, 60], "scale": 1.4},
        {"id": "utility-yard", "kind": "utility_tower", "position": [-80, 135]},
        {"id": "water-tank-a", "kind": "storage_tank", "position": [-135, 110], "scale": 1.1},
        {"id": "water-tank-b", "kind": "storage_tank", "position": [-165, 82], "scale": 0.85},
        {
            "id": "harbor-warehouse-a",
            "kind": "warehouse",
            "position": [-330, -80],
            "heading_degrees": 90,
        },
        {
            "id": "harbor-warehouse-b",
            "kind": "warehouse",
            "position": [-330, -145],
            "heading_degrees": 90,
        },
        {
            "id": "marina-pier-a",
            "kind": "marina_pier",
            "position": [-520, -90],
            "heading_degrees": 90,
        },
        {
            "id": "marina-pier-b",
            "kind": "marina_pier",
            "position": [-520, -190],
            "heading_degrees": 90,
        },
        {"id": "gate-north", "kind": "gatehouse", "position": [40, 245], "heading_degrees": 90},
        {"id": "gate-south", "kind": "gatehouse", "position": [120, -235], "heading_degrees": 270},
    ]
    for index, (east, north) in enumerate(
        [
            (-420, 80),
            (-410, 170),
            (-340, 230),
            (300, 225),
            (390, 150),
            (430, 55),
            (360, -45),
            (300, -145),
        ]
    ):
        assets.append(
            {
                "id": f"urban-block-{index + 1}",
                "kind": "mixed_block",
                "position": [east, north],
                "scale": round(0.8 + (index % 3) * 0.18, 2),
            }
        )
    for index, (east, north) in enumerate([(-270, -60), (-270, -115), (-270, -170), (-210, -170)]):
        assets.append(
            {
                "id": f"cargo-stack-{index + 1}",
                "kind": "cargo_stack",
                "position": [east, north],
                "heading_degrees": 90,
            }
        )
    for index, (east, north, heading) in enumerate(
        [
            (-220, 250, 0),
            (-40, 270, 0),
            (145, 270, 0),
            (270, 205, 90),
            (275, 0, 90),
            (245, -210, 90),
            (40, -265, 180),
            (-185, -245, 180),
        ]
    ):
        assets.append(
            {
                "id": f"perimeter-{index + 1}",
                "kind": "fence_segment",
                "position": [east, north],
                "heading_degrees": heading,
                "scale": 1.6,
            }
        )
    return assets


def _scenario_payload(
    digest: str,
    site_name: str | None,
    sensor_modalities: list[str],
) -> dict[str, Any]:
    modalities = sensor_modalities or ["Radar", "EO", "RF"]
    sensor_positions = [(22, 25), (40, 78), (70, 24), (96, 72), (55, 48), (84, 50)]
    sensors = []
    sensor_ids = []
    used_ids: set[str] = set()
    for index, modality in enumerate(modalities[:6]):
        base = _id(modality)
        sensor_id = base if base not in used_ids else f"{base}-{index + 1}"
        used_ids.add(sensor_id)
        sensor_ids.append(sensor_id)
        sensors.append(
            {
                "id": sensor_id,
                "faction_id": "defense",
                "position": sensor_positions[index],
                "range": 50.0,
                "fov_degrees": 135.0,
                "heading_degrees": float((index * 60 + 15) % 360),
                "capacity": 12,
                "detection_probability": 0.82,
                "measurement_noise": 0.9,
            }
        )
    display_site = site_name or "Coastal Urban Installation"
    return {
        "schema_version": "2.0",
        "scenario_id": f"document-{digest[:12]}",
        "version": 0,
        "name": f"{display_site} — Fictionalized C-UAS"[:120],
        "description": (
            "A document-derived, fictionalized coastal-urban counter-UAS scenario. "
            "It retains broad mission and asset categories while replacing real geometry, "
            "vulnerability locations, and defensive positions with a synthetic layout."
        ),
        "duration_seconds": 150.0,
        "step_seconds": 1.0,
        "seed": int(digest[:8], 16) % 2_147_483_647,
        "evaluation_faction_id": "defense",
        "world": {"width": 120.0, "height": 100.0, "vertical_extent": 10.0},
        "visualization": {"environment": "coastal_urban", "static_assets": _static_assets()},
        "factions": [
            {
                "id": "threat",
                "name": "Surveillance UAS",
                "category": "strategic",
                "color": "#d15c45",
                "objectives": ["Reach a synthetic observation zone"],
            },
            {
                "id": "defense",
                "name": "Installation Defense",
                "category": "strategic",
                "color": "#53a5c8",
                "objectives": ["Detect low-altitude UAS", "Protect synthetic asset categories"],
            },
            {
                "id": "civil",
                "name": "Protected Civil Traffic",
                "category": "background",
                "color": "#c6d3cb",
                "objectives": ["Transit the shared operating area"],
            },
        ],
        "relationships": [
            {
                "source_faction_id": "defense",
                "target_faction_id": "threat",
                "disposition": "opponent",
            },
            {
                "source_faction_id": "defense",
                "target_faction_id": "civil",
                "disposition": "protected",
            },
            {
                "source_faction_id": "threat",
                "target_faction_id": "defense",
                "disposition": "opponent",
            },
            {"source_faction_id": "threat", "target_faction_id": "civil", "disposition": "neutral"},
            {"source_faction_id": "civil", "target_faction_id": "threat", "disposition": "neutral"},
            {
                "source_faction_id": "civil",
                "target_faction_id": "defense",
                "disposition": "neutral",
            },
        ],
        "protected_zones": [
            {"id": "synthetic-operations", "center": [82, 56], "radius": 10.0, "priority": 1.0},
            {"id": "synthetic-utilities", "center": [70, 34], "radius": 8.0, "priority": 0.8},
            {"id": "synthetic-access", "center": [94, 44], "radius": 7.0, "priority": 0.7},
        ],
        "actor_groups": [
            {
                "id": "surveillance-suas",
                "faction_id": "threat",
                "count": 9,
                "mission": "attack_objective",
                "speed": 1.2,
                "start_edge": "north",
                "target_zone_id": "synthetic-operations",
                "attack_probability": 0.0,
                "domain": "air",
            },
            {
                "id": "protected-air-traffic",
                "faction_id": "civil",
                "count": 3,
                "mission": "transit",
                "speed": 0.8,
                "start_edge": "south",
                "domain": "air",
            },
            {
                "id": "protected-maritime-traffic",
                "faction_id": "civil",
                "count": 2,
                "mission": "transit",
                "speed": 0.45,
                "start_edge": "west",
                "domain": "maritime",
            },
        ],
        "sensors": sensors,
        "responses": [
            {
                "id": "response-alpha",
                "faction_id": "defense",
                "position": [58, 42],
                "engagement_range": 34.0,
                "capacity_per_step": 2,
                "effectiveness": 0.58,
            },
            {
                "id": "response-bravo",
                "faction_id": "defense",
                "position": [90, 58],
                "engagement_range": 28.0,
                "capacity_per_step": 1,
                "effectiveness": 0.66,
            },
        ],
        "policy_bindings": [
            {
                "id": "threat-movement",
                "name": "Small UAS surveillance route",
                "faction_id": "threat",
                "role": "movement",
                "policy_id": "probe",
                "implementation": "scripted",
                "trainable": True,
                "actor_group_ids": ["surveillance-suas"],
                "sensor_ids": [],
            },
            {
                "id": "defense-sensing",
                "name": "Layered sensing network",
                "faction_id": "defense",
                "role": "sensing",
                "policy_id": "coverage",
                "implementation": "scripted",
                "trainable": True,
                "actor_group_ids": [],
                "sensor_ids": sensor_ids,
            },
            {
                "id": "civil-movement",
                "name": "Protected air and maritime traffic",
                "faction_id": "civil",
                "role": "movement",
                "policy_id": "transit",
                "implementation": "scripted",
                "trainable": False,
                "actor_group_ids": ["protected-air-traffic", "protected-maritime-traffic"],
                "sensor_ids": [],
            },
        ],
    }


def import_scenario_document(data: bytes, filename: str) -> ScenarioDocumentImport:
    """Convert a JLC-style DOCX background brief into a reviewed scenario draft."""

    embedded_images = _validate_docx(data)
    document = _document_class()(BytesIO(data))
    sections = _extract_sections(document)
    sensor_modalities = _extract_modalities(document)
    response_text = "\n".join(sections.get("responses", []))
    response_modalities = _response_modalities(response_text)
    scope = " ".join(sections.get("scope", []))
    site_name = _extract_site_name(scope)
    digest = hashlib.sha256(data).hexdigest()
    scenario = ScenarioSpec.model_validate(_scenario_payload(digest, site_name, sensor_modalities))
    asset_resolution = [
        AssetResolution(
            source_category="small surveillance UAS",
            scene_asset="quadcopter_surrogate",
            implementation="procedural multirotor with replay truth binding",
        ),
        AssetResolution(
            source_category="layered sensors",
            scene_asset="range_watch_tower",
            implementation="existing detailed asset used as a fused abstract node",
        ),
        AssetResolution(
            source_category="response coordination",
            scene_asset="nasams_fcs",
            implementation="existing detailed asset, labeled only as a generic response vehicle",
        ),
        AssetResolution(
            source_category="protected maritime traffic",
            scene_asset="cargo_ship_liberty",
            implementation="existing detailed traffic surrogate",
        ),
        AssetResolution(
            source_category="coastal and urban context",
            scene_asset="coastal_urban procedural kit",
            implementation=(
                "roads, water, blocks, warehouses, marina, utility yard, and perimeter props"
            ),
        ),
    ]
    report = ScenarioDocumentImportReport(
        source_filename=Path(filename).name,
        source_sha256=digest,
        source_size_bytes=len(data),
        source_site_name=site_name,
        matched_sections=sorted(sections),
        sensor_modalities=sensor_modalities,
        response_modalities=response_modalities,
        embedded_images=embedded_images,
        asset_resolution=asset_resolution,
        safety_transformations=[
            (
                "Replaced source maps, boundaries, and marked points with a fictional "
                "coastal-urban layout."
            ),
            "Converted named infrastructure into generic protected asset categories.",
            "Converted named sensor and effector locations into abstract simulation resources.",
            (
                "Excluded real patrol zones, jurisdiction boundaries, coverage circles, and "
                "approach geometry."
            ),
        ],
        assumptions=[
            "The scenario models optical-surveillance access as a zero-damage objective attempt.",
            (
                "Sensor probabilities, ranges, capacities, and response effectiveness are "
                "synthetic training values."
            ),
            "Protected air and maritime traffic is included to exercise deconfliction.",
        ],
        warnings=[
            "Embedded maps and images were not used as spatial truth.",
            (
                "Authorities and legal statements remain document context and are not encoded "
                "as operational rules."
            ),
            (
                "A human reviewer must approve the scenario before it is used for analysis or "
                "training decisions."
            ),
        ],
    )
    return ScenarioDocumentImport(scenario=scenario, report=report)


def _sensor_label(sensor_id: str) -> str:
    label = sensor_id.replace("-", " ").title()
    return "Fused Sensor Node" if label.startswith("Raven Sensor") else label


def export_scenario_document(
    scenario: ScenarioSpec,
    output_path: Path,
    *,
    hero_image: Path | None = None,
) -> Path:
    """Write a source-like background brief for a scenario."""

    Document = _document_class()
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    document.core_properties.title = f"{scenario.name} Background Data"
    document.core_properties.subject = "Fictional counter-UAS scenario background brief"
    document.core_properties.author = "Pytho Arena"
    document.core_properties.comments = (
        "Fictional and unclassified training surrogate. No real installation geometry or posture."
    )

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    heading = styles["Heading 1"]
    heading.font.name = "Aptos Display"
    heading.font.size = Pt(17)
    heading.font.color.rgb = RGBColor(20, 80, 106)
    heading.font.bold = False
    heading.paragraph_format.space_before = Pt(12)
    heading.paragraph_format.space_after = Pt(4)

    def add_heading(title: str) -> None:
        document.add_heading(title, level=1)

    def add_note(text: str) -> None:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(20, 80, 106)

    def shade(cell: Any, fill: str) -> None:
        properties = cell._tc.get_or_add_tcPr()
        element = OxmlElement("w:shd")
        element.set(qn("w:fill"), fill)
        properties.append(element)

    def style_table(table: Any, header_rows: int = 1, font_size: float = 8.0) -> None:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if row_index < header_rows:
                    shade(cell, "F7E39A")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "Aptos"
                        run.font.size = Pt(font_size)
                        if row_index < header_rows:
                            run.bold = True

    add_note("FICTIONAL / UNCLASSIFIED TRAINING SURROGATE")
    add_heading("Background")
    document.add_paragraph(scenario.description)
    document.add_paragraph(
        "This brief describes a synthetic counter-UAS exercise. It does not describe a real "
        "installation, current force posture, readiness condition, or operational plan."
    )

    add_heading("Scope")
    document.add_paragraph(
        f"The scenario runs for {scenario.duration_seconds:g} seconds at "
        f"{scenario.step_seconds:g}-second decision intervals with seed {scenario.seed}. "
        f"It contains {scenario.entity_count} moving actors, {len(scenario.sensors)} abstract "
        f"sensor nodes, {len(scenario.responses)} generic response elements, and "
        f"{len(scenario.protected_zones)} protected zones."
    )
    document.add_paragraph(
        "Out of scope are real coordinates, routes, unit identities, communications, storage "
        "locations, current defensive positions, readiness, and physical weak points."
    )

    add_heading("Relevant Authorities and Coordination")
    document.add_paragraph(
        "No real command authority, rules of engagement, legal authority, or agency tasking is "
        "encoded. Scenario relationships control eligibility inside the simulator."
    )
    for relationship in scenario.relationships:
        document.add_paragraph(
            f"{relationship.source_faction_id} treats {relationship.target_faction_id} as "
            f"{relationship.disposition}.",
            style="List Bullet",
        )

    add_heading("Base Overview and Protected Assets")
    if scenario.visualization and scenario.visualization.environment == "coastal_urban":
        environment_text = (
            "The fictional scene uses a coastal-urban setting with roads, mixed low-rise blocks, "
            "warehouses, a harbor edge, generic utilities, and a synthetic installation perimeter."
        )
    else:
        environment_text = (
            "The fictional Gulf-desert scene includes one runway, a parallel taxiway, generic "
            "aprons, shelters, hangars, operations and support buildings, a tower, a fire station, "
            "warehouses, tanks, service yards, and abstract counter-UAS props."
        )
    document.add_paragraph(environment_text)
    document.add_paragraph(
        "Scene aircraft and vehicles are visual set dressing unless they also appear in the actor "
        "roster. Geometry and placements are scenario assumptions."
    )
    if hero_image and hero_image.is_file():
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(str(hero_image), width=Inches(6.35))
        caption = document.add_paragraph("Figure 1 - Fictional scenario environment")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.runs[0].italic = True
        caption.runs[0].font.size = Pt(9)

    add_heading("Threat Surveillance Capabilities")
    for group in scenario.actor_groups:
        faction = next(item for item in scenario.factions if item.id == group.faction_id)
        if faction.category != "strategic" or group.faction_id == scenario.evaluation_faction_id:
            continue
        document.add_paragraph(
            f"{group.count} synthetic {group.domain} actors in {group.id.replace('-', ' ')} use a "
            f"{group.mission.replace('_', ' ')} mission. Speed and outcome probability are "
            "normalized simulation inputs, not real-system performance.",
            style="List Bullet",
        )
    document.add_paragraph(
        "Platform attribution, payload, launch point, route geometry, and real-system capability "
        "claims are not part of this scenario."
    )

    add_heading("Blue Surveillance and Tracking")
    document.add_paragraph(
        "Sensor values below are synthetic. Positions and headings are intentionally omitted."
    )
    table = document.add_table(rows=1, cols=6)
    for cell, value in zip(
        table.rows[0].cells,
        ["Sensor", "Capacity", "FOV", "Detection", "Noise", "Range"],
        strict=True,
    ):
        cell.text = value
    for sensor in scenario.sensors:
        cells = table.add_row().cells
        values = [
            _sensor_label(sensor.id),
            str(sensor.capacity),
            f"{sensor.fov_degrees:g}°",
            f"{sensor.detection_probability:.2f}",
            f"{sensor.measurement_noise:g}",
            f"{sensor.range:g} normalized",
        ]
        for cell, value in zip(cells, values, strict=True):
            cell.text = value
    style_table(table, font_size=8.0)
    document.add_paragraph(
        "Operator view shows noisy measurements. Analyst truth remains separate from the "
        "operator track picture."
    )

    document.add_page_break()
    add_heading("Blue Effectors")
    document.add_paragraph(
        "Response elements are generic scoring resources. Rendered vehicles do not identify a "
        "fielded weapon or imply current capability."
    )
    response_table = document.add_table(rows=1, cols=4)
    for cell, value in zip(
        response_table.rows[0].cells,
        ["Response", "Capacity/step", "Effectiveness", "Range"],
        strict=True,
    ):
        cell.text = value
    for response in scenario.responses:
        cells = response_table.add_row().cells
        values = [
            response.id.replace("-", " ").title(),
            str(response.capacity_per_step),
            f"{response.effectiveness:.2f}",
            f"{response.engagement_range:g} normalized",
        ]
        for cell, value in zip(cells, values, strict=True):
            cell.text = value
    style_table(response_table, font_size=8.0)
    document.add_paragraph(
        "Weapon type, emplacement, inventory, reload, readiness, response timing, and real "
        "engagement ranges are not modeled."
    )

    add_heading("Detection and Synthetic Performance Views")
    document.add_paragraph(
        "Operator view presents detections and current events while withholding active truth and "
        "aggregate outcomes. Analyst view adds replay truth, entity focus, detection and defense "
        "rates, intercepts, objective attempts, and protected-traffic status."
    )
    document.add_paragraph(
        "Results may identify synthetic missed detections, capacity saturation, unsuccessful "
        "responses, or protected-traffic effects. They do not identify a real coverage gap, "
        "attack route, installation weakness, or readiness condition."
    )

    add_heading("References")
    references = [
        "Pytho Arena scenario JSON and immutable content hash.",
        "Pytho Arena operator/analyst replay and relationship behavior.",
        "Pytho Arena docs/SCENARIO_BASIS.md (scenario and public visual-reference basis).",
        "U.S. Central Command, public releases cited in SCENARIO_BASIS.md (regional context only).",
        (
            "U.S. Air Forces Central, DVIDS, and USACE public imagery cited in "
            "SCENARIO_BASIS.md (visual context only)."
        ),
    ]
    for reference in references:
        document.add_paragraph(reference, style="List Number")

    add_heading("Traceability")
    trace = document.add_table(rows=1, cols=4)
    for cell, value in zip(
        trace.rows[0].cells,
        ["Statement", "Source", "Class", "Limitation"],
        strict=True,
    ):
        cell.text = value
    rows = [
        ("Actor, sensor, response counts", "Scenario JSON", "Repository fact", "Synthetic inputs"),
        (
            "Environment and static assets",
            "Visualization spec",
            "Scenario assumption",
            "Fictional geometry",
        ),
        ("Operator/analyst views", "Arena UI contract", "Repository fact", "Exercise display only"),
        ("Regional visual context", "SCENARIO_BASIS.md", "Public context", "No real layout copied"),
    ]
    for values in rows:
        cells = trace.add_row().cells
        for cell, value in zip(cells, values, strict=True):
            cell.text = value
    style_table(trace, font_size=7.5)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path
